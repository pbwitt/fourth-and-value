"""One-off source probe for the low-scoring study: prints coverage and field names, writes nothing."""
import io
import json
import requests

S = requests.Session()
S.headers['User-Agent'] = 'Mozilla/5.0 (research; fourthandvalue.com)'


def get(url, **kw):
    try:
        r = S.get(url, timeout=60, **kw)
        print('GET', r.status_code, r.url[:200], len(r.content))
        return r
    except Exception as e:
        print('ERR', url, e)


print('== CDX mlb')
for pat in ['sportsbookreviewsonline.com/scoresoddsarchives/mlb/*', 'sportsbookreviewsonline.com/scoresoddsarchives/mlb*']:
    r = get('http://web.archive.org/cdx/search/cdx', params=dict(url=pat, output='json', collapse='urlkey', limit=500))
    if r is not None and r.ok:
        for row in r.json()[1:]:
            print(row[1], row[4], row[2])

print('== SBRO samples')
import openpyxl
for u in ['https://sportsbookreviewsonline.com/scoresoddsarchives/nhl/nhl%20odds%202015-16.xlsx',
          'https://sportsbookreviewsonline.com/scoresoddsarchives/nba/nba%20odds%202015-16.xlsx',
          'https://sportsbookreviewsonline.com/scoresoddsarchives/nhl/nhl%20odds%202022-23.xlsx',
          'https://sportsbookreviewsonline.com/scoresoddsarchives/mlb/mlb%20odds%202015.xlsx']:
    r = get('http://web.archive.org/web/2023id_/' + u)
    if r is None or not r.ok or r.content[:2] != b'PK':
        print(' not an xlsx', r.content[:80] if r is not None else None)
        continue
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    ws = wb.worksheets[0]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 7 or i in (500, 501):
            print(' ', row)
        if i > 501:
            break
    print(' rows', ws.max_row)

print('== ESPN core odds coverage by season')
for sport, lg, dates in [('hockey', 'nhl', ['20090115', '20120115', '20150115', '20170115', '20180115', '20200115', '20220115', '20240115']),
                         ('basketball', 'nba', ['20090115', '20120115', '20150115', '20170115', '20180115', '20220115', '20240115']),
                         ('baseball', 'mlb', ['20100715', '20140715', '20170715', '20180715', '20190715', '20210715', '20230715', '20250715'])]:
    for d in dates:
        r = get(f'https://site.api.espn.com/apis/site/v2/sports/{sport}/{lg}/scoreboard', params=dict(dates=d, limit=50))
        if r is None or not r.ok or not r.json().get('events'):
            continue
        eid = r.json()['events'][0]['id']
        o = get(f'https://sports.core.api.espn.com/v2/sports/{sport}/leagues/{lg}/events/{eid}/competitions/{eid}/odds')
        if o is not None and o.ok:
            items = o.json().get('items', [])
            print(lg, d, 'providers', [(i.get('provider', {}).get('name'), i.get('overUnder')) for i in items][:6])
            if items and d in ('20240115', '20180715', '20180115'):
                it = items[0]
                print('  item keys', list(it.keys()))
                print('  total fields', {k: it.get(k) for k in ('overUnder', 'overOdds', 'underOdds')}, json.dumps({k: it.get(k) for k in ('open', 'close', 'current')})[:800])
        f = get(f'https://sports.core.api.espn.com/v2/sports/{sport}/leagues/{lg}/events/{eid}/competitions/{eid}/officials')
        if f is not None and f.ok:
            print('  officials', json.dumps(f.json())[:600])

print('== MLB schedule hydrate')
for start, end in [('2010-07-01', '2010-07-02'), ('2016-07-01', '2016-07-02'), ('2021-07-01', '2021-07-02'), ('2025-07-01', '2025-07-02')]:
    r = get('https://statsapi.mlb.com/api/v1/schedule', params=dict(sportId=1, startDate=start, endDate=end, hydrate='weather,officials,venue,linescore,game(content(summary))'))
    if r is not None and r.ok:
        g = r.json()['dates'][0]['games'][0]
        print(' keys', list(g.keys()))
        print(' weather', g.get('weather'), 'officials', json.dumps(g.get('officials'))[:300], 'dayNight', g.get('dayNight'), 'dh', g.get('doubleHeader'), g.get('gameDate'))
    r = get('https://statsapi.mlb.com/api/v1/schedule', params=dict(sportId=1, startDate=start, endDate=end))
    if r is not None and r.ok:
        pk = r.json()['dates'][0]['games'][0]['gamePk']
        f = get(f'https://statsapi.mlb.com/api/v1.1/game/{pk}/feed/live')
        if f is not None and f.ok:
            j = f.json()
            print(' feed weather', j['gameData'].get('weather'), 'datetime', j['gameData'].get('datetime'), 'officials', json.dumps(j['liveData']['boxscore'].get('officials'))[:300])
            print(' gameInfo', json.dumps(j['liveData']['boxscore'].get('info'))[:800])

print('== NHL right-rail 2008')
r = get('https://api-web.nhle.com/v1/gamecenter/2008020500/right-rail')
if r is not None and r.ok:
    print(json.dumps(r.json().get('gameInfo'))[:400])

print('== nflverse pbp weather')
r = get('https://github.com/nflverse/nflverse-data/releases/download/pbp/play_by_play_2015.csv.gz', stream=True)
if r is not None:
    print(' ok', r.ok)
